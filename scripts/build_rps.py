#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
build_rps.py - Generator / Formatter Dokumen RPS UNIROW OBE 2026
Membuat file .docx RPS yang 100% presisi dengan format resmi Universitas PGRI Ronggolawe (UNIROW) Tuban
berdasarkan dokumen acuan: 01_1_RPS_Antropologi Budaya_OBE.docx
"""

import os
import sys
import json
import re
import argparse
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
SKILL_DIR = os.path.dirname(SCRIPT_DIR)
DEFAULT_TEMPLATE = os.path.join(SKILL_DIR, 'assets', 'Template_RPS_UNIROW.docx')

def add_markdown_runs(paragraph, text, font_size=10.5, base_bold=False, base_italic=False, color_rgb=None):
    """
    Menambahkan run ke paragraph dengan mem-parsing token markdown inline:
    - **teks tebal** -> run.bold = True
    - *teks miring* -> run.italic = True
    - ***teks tebal miring*** -> run.bold = True, run.italic = True
    Menghilangkan seluruh karakter bintang asterisks literal.
    """
    if not text:
        return

    pattern = re.compile(r'(\*\*\*.*?\*\*\*|\*\*.*?\*\*|\*[^*\n]+?\*)')
    tokens = pattern.split(str(text))
    for token in tokens:
        if not token:
            continue
        bold = base_bold
        italic = base_italic
        clean = token
        if token.startswith('***') and token.endswith('***') and len(token) >= 6:
            bold = True
            italic = True
            clean = token[3:-3]
        elif token.startswith('**') and token.endswith('**') and len(token) >= 4:
            bold = True
            clean = token[2:-2]
        elif token.startswith('*') and token.endswith('*') and len(token) >= 2:
            italic = True
            clean = token[1:-1]
        
        clean = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', clean.replace('**', '').replace('__', ''))
        if clean:
            run = paragraph.add_run(clean)
            run.bold = bold
            run.italic = italic
            run.font.name = 'Cambria'
            run.font.size = Pt(font_size)
            if color_rgb:
                run.font.color.rgb = color_rgb

def set_cell_text(cell, text, font_size=10.5, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.LEFT, color_rgb=None):
    cell.text = ''
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    
    clean_full_text = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', str(text))
    lines = clean_full_text.split('\n')
    for idx, line in enumerate(lines):
        if idx > 0:
            p = cell.add_paragraph()
            p.alignment = align
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
        
        if bold:
            clean_line = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', line.replace('**', '').replace('*', '').replace('__', ''))
            run = p.add_run(clean_line)
            run.bold = True
            run.italic = italic
            run.font.name = 'Cambria'
            run.font.size = Pt(font_size)
            if color_rgb:
                run.font.color.rgb = color_rgb
        else:
            add_markdown_runs(p, line, font_size=font_size, base_bold=False, base_italic=italic, color_rgb=color_rgb)

def set_cell_padding(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for edge, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{edge}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_shading(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_table_borders(table, color="B0B5B3", sz="4"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:right w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def build_nested_matrix(cell, cpl_list, matrix_rows):
    """
    Membangun / memperbarui matriks korelasi CPL vs Sub-CPMK (nested table di Row 13).
    Disesuaikan dinamis dengan jumlah CPL pada mata kuliah.
    """
    if cell.tables:
        tbl_elem = cell.tables[0]._tbl
        tbl_elem.getparent().remove(tbl_elem)
        
    num_cpls = len(cpl_list)
    num_rows = len(matrix_rows) + 1
    num_cols = num_cpls + 2
    
    new_tbl = cell.add_table(rows=num_rows, cols=num_cols)
    set_table_borders(new_tbl, color="B0B5B3", sz="4")
    
    # Calculate widths (total ~ 13400 dxa)
    total_w = 13400
    col0_w = 2600
    col_last_w = 2000
    cpl_w = int((total_w - col0_w - col_last_w) / max(1, num_cpls))
    
    # Header
    hdr = new_tbl.rows[0]
    set_cell_text(hdr.cells[0], 'Sub-CPMK / Evaluasi', font_size=10.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_shading(hdr.cells[0], 'EAECEE')
    set_cell_padding(hdr.cells[0])
    
    for i, cpl in enumerate(cpl_list):
        cpl_title = f"{cpl} (%)" if "(%)" not in cpl else cpl
        set_cell_text(hdr.cells[i + 1], cpl_title, font_size=10.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(hdr.cells[i + 1], 'EAECEE')
        set_cell_padding(hdr.cells[i + 1])
        
    set_cell_text(hdr.cells[-1], 'Bobot Penilaian (%)', font_size=10.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_shading(hdr.cells[-1], 'EAECEE')
    set_cell_padding(hdr.cells[-1])
    
    # Data rows
    for r_idx, r_data in enumerate(matrix_rows):
        row = new_tbl.rows[r_idx + 1]
        lbl = str(r_data.get('label', '')).replace('**', '').replace('*', '').strip()
        is_total = (r_idx == len(matrix_rows) - 1) or 'TOTAL' in lbl.upper()
        is_exam = lbl in ['UTS', 'UAS'] or 'UJIAN' in lbl.upper()
        
        bold_flag = is_total or is_exam
        bg_fill = 'EAECEE' if is_total else ('F4F6F7' if is_exam else None)
        
        # Col 0: Sub-CPMK / Evaluasi label
        set_cell_text(row.cells[0], lbl, font_size=10.5, bold=bold_flag, align=WD_ALIGN_PARAGRAPH.CENTER if is_exam or is_total else WD_ALIGN_PARAGRAPH.LEFT)
        if bg_fill:
            set_cell_shading(row.cells[0], bg_fill)
        set_cell_padding(row.cells[0])
        
        # CPL checkmarks
        cpls_dict = r_data.get('cpls', {})
        for i, cpl in enumerate(cpl_list):
            val = str(cpls_dict.get(cpl, '')).replace('**', '').replace('*', '').strip()
            set_cell_text(row.cells[i + 1], val, font_size=10.5, bold=bold_flag, align=WD_ALIGN_PARAGRAPH.CENTER)
            if bg_fill:
                set_cell_shading(row.cells[i + 1], bg_fill)
            set_cell_padding(row.cells[i + 1])
            
        # Bobot Penilaian col
        bobot_val = str(r_data.get('bobot', '')).replace('**', '').replace('*', '').strip()
        set_cell_text(row.cells[-1], bobot_val, font_size=10.5, bold=bold_flag, align=WD_ALIGN_PARAGRAPH.CENTER)
        if bg_fill:
            set_cell_shading(row.cells[-1], bg_fill)
        set_cell_padding(row.cells[-1])
        
    return new_tbl

def extract_soal_uas(text):
    """
    Mengekstrak naskah soal UAS dari format Markdown, mendukung naskah soal standar HOTS
    maupun naskah soal berbasis Kasus Posisi (Hukum/Studi Kasus).
    """
    soal_block = re.search(r'##\s*(?:CONTOH\s+)?SOAL\s+UJIAN\s+AKHIR\s+SEMESTER.*?\n(.*?)(?=\n##\s+[A-Z]|\Z)', text, re.DOTALL | re.IGNORECASE)
    if not soal_block:
        return []
    sb = soal_block.group(1)
    # Bersihkan petunjuk pengerjaan umum
    sb = re.sub(r'###\s*Petunjuk Pengerjaan:.*?(?=###|\Z)', '', sb, flags=re.DOTALL | re.IGNORECASE)
    
    # Cek apakah format Kasus Posisi
    if re.search(r'###\s*Kasus Posisi', sb, re.IGNORECASE):
        cases = re.split(r'###\s*Kasus Posisi\s*', sb, flags=re.IGNORECASE)[1:]
        soal_list = []
        for c in cases:
            lines = c.strip().split('\n')
            case_title = lines[0].strip().replace('**', '').replace('*', '').strip()
            sub_q = [l.strip() for l in lines if re.match(r'^\d+\.\s+', l.strip())]
            for sq in sub_q:
                sq_clean = re.sub(r'^\d+\.\s*', '', sq).strip()
                soal_list.append(f"[{case_title}] {sq_clean}")
        return soal_list
    else:
        # Format langsung nomor 1, 2, 3...
        items = re.findall(r'(?:^|\n)(\d+\.\s+[^\n]+(?:\n(?!\d+\.)[^\n]+)*)', sb)
        clean_items = []
        for it in items:
            it_clean = ' '.join([l.strip() for l in it.split('\n') if l.strip()])
            it_clean = re.sub(r'^\d+\.\s*', '', it_clean).strip()
            clean_items.append(it_clean)
        return clean_items

def parse_rps_markdown(md_path):
    """
    Mem-parsing berkas RPS berformat Markdown menjadi dictionary data lengkap.
    """
    with open(md_path, 'r', encoding='utf-8') as f:
        text = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', f.read())

    data = {
        'identitas': {},
        'otorisasi': {},
        'capaian_pembelajaran': {
            'cpl_prodi': [],
            'cpmk': [],
            'sub_cpmk': [],
            'matriks_korelasi': []
        },
        'detail_mk': {},
        'mingguan': [],
        'rubrik': [],
        'evaluasi_kelulusan': {},
        'soal_uas': [],
        'validasi': {}
    }

    # 1. Identitas MK
    id_block = re.search(r'## IDENTITAS MATA KULIAH\s*(.*?)(?=##|\Z)', text, re.DOTALL)
    if id_block:
        for line in id_block.group(1).split('\n'):
            line = line.strip()
            if '|' in line and not line.startswith('|---|'):
                parts = [p.strip() for p in line.split('|')[1:-1]]
                if len(parts) >= 2:
                    k = parts[0].replace('**', '').strip()
                    v = parts[1].replace('**', '').strip()
                    if 'MATA KULIAH' in k.upper():
                        data['identitas']['nama_mk'] = v
                    elif 'KODE' in k.upper():
                        data['identitas']['kode_mk'] = v
                    elif 'RUMPUN' in k.upper():
                        data['identitas']['rumpun_mk'] = v
                    elif 'BOBOT' in k.upper():
                        sks_num = re.search(r'\d+', v)
                        data['identitas']['bobot_sks'] = int(sks_num.group()) if sks_num else 2
                    elif 'SEMESTER' in k.upper():
                        sem_num = re.search(r'\d+', v)
                        if not sem_num:
                            if 'III' in v.upper():
                                data['identitas']['semester'] = 3
                            elif 'II' in v.upper():
                                data['identitas']['semester'] = 2
                            else:
                                data['identitas']['semester'] = 1
                        else:
                            data['identitas']['semester'] = int(sem_num.group())
                    elif 'TGL' in k.upper():
                        data['identitas']['tgl_penyusunan'] = v

    sem_val = data['identitas'].get('semester', 1)
    romans = {1: 'I', 2: 'II', 3: 'III', 4: 'IV', 5: 'V', 6: 'VI', 7: 'VII', 8: 'VIII'}
    sem_roman = romans.get(sem_val, 'I')
    data['identitas']['fakultas'] = 'Keguruan dan Ilmu Pendidikan'
    data['identitas']['prodi'] = 'Pendidikan Pancasila dan Kewarganegaraan'
    data['identitas']['kode_dokumen'] = f'PPKn/{sem_roman}/PKN/MKK PPKn'

    # 2. Otorisasi
    otor_block = re.search(r'## OTORISASI / PENGESAHAN\s*(.*?)(?=##|\Z)', text, re.DOTALL)
    if otor_block:
        lines = [l.strip() for l in otor_block.group(1).split('\n') if l.strip() and '|' in l and not re.match(r'^[|\s:-]+$', l)]
        if len(lines) >= 2:
            names = [c.replace('**', '').replace('<br>', ' ').strip() for c in lines[-1].split('|')[1:-1]]
            clean_names = []
            for n in names:
                n_clean = n.split('NIDN')[0].strip()
                clean_names.append(n_clean)
            if len(clean_names) >= 3:
                data['otorisasi']['pengembang'] = clean_names[0]
                data['otorisasi']['koordinator_rmk'] = clean_names[1]
                data['otorisasi']['kaprodi'] = clean_names[2]

    # 3. Capaian Pembelajaran (CP)
    # CPL: match - **CPL...**: ...
    cpl_matches = re.findall(r'-\s*\*\*(CPL[-_\s]*\d+[^:]*?)\*\*:\s*(.+)', text)
    for kode_raw, deskripsi in cpl_matches:
        num_m = re.search(r'\d+', kode_raw)
        cpl_norm = f"CPL{num_m.group()}" if num_m else kode_raw
        data['capaian_pembelajaran']['cpl_prodi'].append({
            'kode': cpl_norm,
            'deskripsi': deskripsi.strip()
        })

    # CPMK: match - **CPMK...**: ...
    cpmk_matches = re.findall(r'-\s*\*\*(CPMK[-_\s]*\d+[^:]*?)\*\*:\s*(.+)', text)
    for kode_raw, deskripsi in cpmk_matches:
        num_m = re.search(r'\d+', kode_raw)
        cpmk_norm = f"CPMK{num_m.group()}" if num_m else kode_raw
        data['capaian_pembelajaran']['cpmk'].append({
            'kode': cpmk_norm,
            'deskripsi': deskripsi.strip()
        })

    # Sub-CPMK: match - **Sub-CPMK...** [Takso]: ...
    sub_matches = re.findall(r'-\s*\*\*(Sub-CPMK[-_\s]*\d+[^:]*?)\*\*\s*(?:\[(.*?)\])?:\s*(.+)', text)
    for kode_raw, takso, deskripsi in sub_matches:
        num_m = re.search(r'\d+', kode_raw)
        sub_norm = f"Sub-CPMK {num_m.group()}" if num_m else kode_raw
        data['capaian_pembelajaran']['sub_cpmk'].append({
            'kode': sub_norm,
            'taksonomi': (takso or '').strip(),
            'deskripsi': deskripsi.strip()
        })

    # 4. Matriks Korelasi
    mat_block = re.search(r'## MATRIKS KORELASI.*?\n(.*?)(?=##|\Z)', text, re.DOTALL)
    if mat_block:
        lines = [l.strip() for l in mat_block.group(1).split('\n') if l.strip() and '|' in l and not re.match(r'^[|\s:-]+$', l)]
        if len(lines) >= 2:
            headers = [c.strip() for c in lines[0].split('|')[1:-1]]
            cpl_hdr = []
            for h in headers[1:-1]:
                h_clean = h.replace('(%)', '').strip()
                cpl_hdr.append(h_clean)
            
            for l in lines[1:]:
                cols = [c.strip() for c in l.split('|')[1:-1]]
                if cols:
                    lbl = cols[0].replace('**', '').strip()
                    bobot = cols[-1].replace('**', '').strip()
                    cpls_dict = {}
                    for ci, c_code in enumerate(cpl_hdr):
                        if ci + 1 < len(cols) - 1:
                            val = cols[ci + 1].strip()
                            if val:
                                cpls_dict[c_code] = val
                    data['capaian_pembelajaran']['matriks_korelasi'].append({
                        'label': lbl,
                        'cpls': cpls_dict,
                        'bobot': bobot
                    })

    # 5. Detail MK
    det_block = re.search(r'## DETAIL MATA KULIAH\s*(.*?)(?=\n##\s+[A-Z]|\Z)', text, re.DOTALL)
    if det_block:
        dt = det_block.group(1)
        desk_m = re.search(r'(?:###?\s*Deskripsi Singkat[^\n]*|\*\*Deskripsi Singkat[^\n*]*\*\*)\s*\n*(.+?)(?=(?:###?\s*Bahan Kajian|\*\*Bahan Kajian)|\Z)', dt, re.DOTALL | re.IGNORECASE)
        if desk_m:
            data['detail_mk']['deskripsi'] = desk_m.group(1).strip()
            
        bk_m = re.search(r'(?:###?\s*Bahan Kajian[^\n]*|\*\*Bahan Kajian[^\n*]*\*\*)\s*\n*(.+?)(?=(?:###?\s*Pustaka|\*\*Pustaka)|\Z)', dt, re.DOTALL | re.IGNORECASE)
        if bk_m:
            data['detail_mk']['bahan_kajian'] = bk_m.group(1).strip()

        pu_m = re.search(r'(?:####?\s*(?:Pustaka\s+)?Utama[^\n]*|\*\*(?:Pustaka\s+)?Utama[^\n*]*\*\*)\s*\n*(.+?)(?=(?:####?\s*(?:Pustaka\s+)?Pendukung|\*\*(?:Pustaka\s+)?Pendukung)|\Z)', dt, re.DOTALL | re.IGNORECASE)
        if pu_m:
            pu_lines = [l.strip() for l in pu_m.group(1).split('\n') if l.strip() and not l.strip().startswith('#')]
            data['detail_mk']['pustaka_utama'] = pu_lines

        pp_m = re.search(r'(?:####?\s*(?:Pustaka\s+)?Pendukung[^\n]*|\*\*(?:Pustaka\s+)?Pendukung[^\n*]*\*\*)\s*\n*(.+?)(?=(?:###?\s*Dosen Pengampu|\*\*Dosen Pengampu)|\Z)', dt, re.DOTALL | re.IGNORECASE)
        if pp_m:
            pp_lines = [l.strip() for l in pp_m.group(1).split('\n') if l.strip() and not l.strip().startswith('#')]
            data['detail_mk']['pustaka_pendukung'] = pp_lines

        dp_m = re.search(r'(?:###?\s*Dosen Pengampu[^\n]*|\*\*Dosen Pengampu[^\n*]*\*\*)\s*\n*(.+?)(?=(?:###?\s*Mata\s*kuliah Syarat|\*\*Mata\s*kuliah Syarat)|\Z)', dt, re.DOTALL | re.IGNORECASE)
        if dp_m:
            dp_lines = [l.strip().lstrip('-* ').strip() for l in dp_m.group(1).split('\n') if l.strip()]
            data['detail_mk']['dosen_pengampu'] = ', '.join(dp_lines)

        ms_m = re.search(r'(?:###?\s*Mata\s*kuliah Syarat[^\n]*|\*\*Mata\s*kuliah Syarat[^\n*]*\*\*)\s*\n*(.+?)(?=\n\n|\Z|---)', dt, re.DOTALL | re.IGNORECASE)
        if ms_m:
            data['detail_mk']['matakuliah_syarat'] = ms_m.group(1).strip().lstrip('-* ').strip()

    # 6. Rencana Pembelajaran 16 Minggu
    mg_block = re.search(r'##\s*RENCANA\s+(?:KEGIATAN\s+PEMBELAJARAN|PEMBELAJARAN\s+MINGGUAN).*?\n(.*?)(?=\n##\s+[A-Z]|\Z)', text, re.DOTALL | re.IGNORECASE)
    if mg_block:
        lines = [l.strip() for l in mg_block.group(1).split('\n') if l.strip() and '|' in l and not re.match(r'^[|\s:-]+$', l)]
        for l in lines:
            cols = [c.strip() for c in l.split('|')[1:-1]]
            if not cols:
                continue
            # skip header rows like (1) or Mg Ke-
            col0_clean = cols[0].replace('**', '').replace('*', '').strip()
            if col0_clean.startswith('(') or 'MG KE' in col0_clean.upper() or not col0_clean:
                continue
            mg_match = re.search(r'\b(\d{1,2})\b', col0_clean)
            if not mg_match:
                if re.search(r'\bUTS\b', col0_clean.upper()) or 'TENGAH SEMESTER' in col0_clean.upper():
                    mg_no = 8
                elif re.search(r'\bUAS\b', col0_clean.upper()) or 'AKHIR SEMESTER' in col0_clean.upper():
                    mg_no = 16
                else:
                    continue
            else:
                mg_no = int(mg_match.group())

            if mg_no < 1 or mg_no > 16:
                continue

            bobot_str = cols[-1].replace('**', '').replace('%', '').strip()
            bobot_val = int(bobot_str) if bobot_str.isdigit() else 0

            # Sesuai kurikulum OBE UNIROW, UTS mutlak di Mg 8 dan UAS mutlak di Mg 16.
            # Hindari substring 'UAS' dalam kata seperti 'mengevaluasi', 'penguasaan', 'luas'.
            is_uts_row = (mg_no == 8)
            is_uas_row = (mg_no == 16)

            # 8-column format (Semester 1 style)
            if len(cols) >= 8:
                if is_uts_row:
                    data['mingguan'].append({
                        'minggu': 8,
                        'is_uts': True,
                        'indikator': cols[2].replace('<br>', '\n').strip() if len(cols) > 2 and cols[2] else 'Evaluasi penguasaan materi perkuliahan minggu 1 s.d. 7',
                        'bobot': bobot_val or 25
                    })
                elif is_uas_row:
                    data['mingguan'].append({
                        'minggu': 16,
                        'is_uas': True,
                        'indikator': cols[2].replace('<br>', '\n').strip() if len(cols) > 2 and cols[2] else 'Evaluasi komprehensif penguasaan capaian pembelajaran mata kuliah',
                        'bobot': bobot_val or 25
                    })
                else:
                    data['mingguan'].append({
                        'minggu': mg_no,
                        'sub_cpmk': cols[1].replace('**', '').strip(),
                        'indikator': cols[2].replace('<br>', '\n').strip(),
                        'teknik_kriteria': cols[3].replace('<br>', '\n').strip(),
                        'luring': cols[4].replace('<br>', '\n').strip(),
                        'daring': cols[5].replace('<br>', '\n').strip(),
                        'materi_pustaka': cols[6].replace('<br>', '\n').strip(),
                        'bobot': bobot_val
                    })
            # 7-column format (Semester 3 style)
            elif len(cols) == 7:
                if is_uts_row:
                    data['mingguan'].append({
                        'minggu': 8,
                        'is_uts': True,
                        'indikator': cols[3].replace('<br>', ' ').strip() if len(cols) > 3 and cols[3] else 'Evaluasi penguasaan materi perkuliahan minggu 1 s.d. 7',
                        'bobot': bobot_val or 25
                    })
                elif is_uas_row:
                    data['mingguan'].append({
                        'minggu': 16,
                        'is_uas': True,
                        'indikator': cols[3].replace('<br>', ' ').strip() if len(cols) > 3 and cols[3] else 'Evaluasi komprehensif penguasaan capaian pembelajaran mata kuliah',
                        'bobot': bobot_val or 25
                    })
                else:
                    sub_txt = cols[1].replace('**', '').strip()
                    bentuk = cols[2].replace('<br>', '\n').strip()
                    ind = cols[3].replace('<br>', '\n').strip()
                    tek = cols[4].replace('<br>', '\n').strip()
                    materi = cols[5].replace('<br>', '\n').strip()
                    
                    lur = bentuk
                    dar = '-'
                    if '**Luring:**' in bentuk or 'Luring:' in bentuk:
                        parts = re.split(r'\*{0,2}Daring:\*{0,2}', bentuk, flags=re.IGNORECASE)
                        if len(parts) > 1:
                            lur = re.sub(r'^\*{0,2}Luring:\*{0,2}\s*', '', parts[0], flags=re.IGNORECASE).strip()
                            dar = parts[1].strip()
                    
                    data['mingguan'].append({
                        'minggu': mg_no,
                        'sub_cpmk': sub_txt,
                        'indikator': ind,
                        'teknik_kriteria': tek,
                        'luring': lur,
                        'daring': dar or '-',
                        'materi_pustaka': materi,
                        'bobot': bobot_val
                    })
            # 5 or 6-column format (Semester 2 style)
            elif len(cols) >= 5:
                if is_uts_row:
                    data['mingguan'].append({
                        'minggu': 8,
                        'is_uts': True,
                        'indikator': cols[2] if len(cols) > 2 and cols[2] else 'Evaluasi penguasaan materi perkuliahan minggu 1 s.d. 7',
                        'bobot': bobot_val or 24
                    })
                elif is_uas_row:
                    data['mingguan'].append({
                        'minggu': 16,
                        'is_uas': True,
                        'indikator': cols[2] if len(cols) > 2 and cols[2] else 'Evaluasi komprehensif penguasaan capaian pembelajaran mata kuliah',
                        'bobot': bobot_val or 26
                    })
                else:
                    sub_txt = cols[1].replace('**', '').strip()
                    penilaian = cols[2] if len(cols) > 2 else ''
                    bentuk = cols[3] if len(cols) > 3 else ''
                    materi = cols[4] if len(cols) > 4 else ''
                    
                    ind = ''
                    tek = ''
                    if '**Indikator:**' in penilaian:
                        p_parts = penilaian.split('**Indikator:**')[1]
                        if '<br>**Teknik:**' in p_parts:
                            ind = p_parts.split('<br>**Teknik:**')[0].strip()
                            tek_crit = p_parts.split('<br>**Teknik:**')[1].strip()
                            tek = 'Teknik: ' + tek_crit.replace('<br>**Kriteria:**', '\nKriteria:').strip()
                        else:
                            ind = p_parts.strip()
                    else:
                        ind = penilaian
                        
                    lur = ''
                    dar = ''
                    if '**Luring:**' in bentuk:
                        b_parts = bentuk.split('**Luring:**')[1]
                        if '<br>**Daring:**' in b_parts:
                            lur = b_parts.split('<br>**Daring:**')[0].strip()
                            dar = b_parts.split('<br>**Daring:**')[1].strip()
                        else:
                            lur = b_parts.strip()
                    else:
                        lur = bentuk

                    data['mingguan'].append({
                        'minggu': mg_no,
                        'sub_cpmk': sub_txt,
                        'indikator': ind,
                        'teknik_kriteria': tek,
                        'luring': lur,
                        'daring': dar or '-',
                        'materi_pustaka': materi,
                        'bobot': bobot_val
                    })

    # 7. Rubrik Penilaian
    lines = text.split('\n')
    in_rub = False
    for l in lines:
        if 'Aspek Penilaian' in l and '|' in l:
            in_rub = True
            continue
        if in_rub:
            if not l.strip() or not l.startswith('|') or l.startswith('###') or l.startswith('##'):
                in_rub = False
                continue
            if re.match(r'^[|\s:-]+$', l):
                continue
            cols = [c.strip() for c in l.split('|')[1:-1]]
            if len(cols) >= 4:
                aspek_raw = cols[0].replace('**', '').strip()
                bobot_m = re.search(r'\((\d+)%\)', aspek_raw)
                aspek_clean = re.sub(r'\s*\(\d+%\)', '', aspek_raw).strip()
                
                if len(cols) >= 6:
                    bobot_val = cols[-1].replace('%', '').strip()
                elif len(cols) == 5:
                    bobot_val = cols[-1].replace('%', '').strip() if re.match(r'^\d+$', cols[-1].replace('%', '').strip()) else (bobot_m.group(1) if bobot_m else '25')
                else:
                    bobot_val = bobot_m.group(1) if bobot_m else '25'
                    
                data['rubrik'].append({
                    'aspek': aspek_clean,
                    'bobot': bobot_val,
                    'sangat_baik': cols[1].strip(),
                    'baik': cols[2].strip(),
                    'cukup': cols[3].strip()
                })

    # 7.5 Evaluasi Kelulusan
    eval_block = re.search(r'(?:Kriteria Kelulusan|Komponen Penilaian Hasil Belajar).*?\n(.*?)(?=##|\Z)', text, re.DOTALL | re.IGNORECASE)
    if eval_block:
        for line in eval_block.group(1).split('\n'):
            line = line.strip()
            if '|' in line and not re.match(r'^[|\s:-]+$', line):
                parts = [p.strip() for p in line.split('|')[1:-1]]
                if len(parts) >= 2:
                    k, v = parts[0].lower(), re.search(r'\d+', parts[1])
                    val = int(v.group()) if v else 0
                    if 'keaktifan' in k or 'kehadiran' in k:
                        data['evaluasi_kelulusan']['keaktifan'] = val
                    elif 'tugas' in k or 'kuis' in k:
                        data['evaluasi_kelulusan']['tugas'] = val
                    elif 'uts' in k:
                        data['evaluasi_kelulusan']['uts'] = val
                    elif 'proyek' in k:
                        data['evaluasi_kelulusan']['proyek'] = val
                    elif 'uas' in k:
                        data['evaluasi_kelulusan']['uas'] = val

    # 8. Validasi
    val_block = re.search(r'##\s*LEMBAR\s+(?:VALIDASI|PENGESAHAN).*?\n(.*?)(?=\n##\s+[A-Z]|\Z)', text, re.DOTALL | re.IGNORECASE)
    if val_block:
        vt = val_block.group(1)
        tgl_m = re.search(r'(?:Divalidasi di Tuban, pada tanggal:|telah divalidasi pada tanggal:?|Tuban,)\s*([^\n\r]+)', vt, re.IGNORECASE)
        if tgl_m:
            data['validasi']['tgl_validasi'] = tgl_m.group(1).replace('**', '').strip()
        
        lines = [l.strip() for l in vt.split('\n') if l.strip() and '|' in l and not re.match(r'^[|\s:-]+$', l)]
        if len(lines) >= 2:
            row_idx = -2 if len(lines) >= 3 and not any(kw in lines[-2].lower() for kw in ['menyetujui', 'mengetahui']) else -1
            row_names = [c.replace('**', '').replace('<br>', ' ').strip() for c in lines[row_idx].split('|')[1:-1]]
            clean_rnames = [n.split('NIDN')[0].strip() for n in row_names if n.strip()]
            if len(clean_rnames) >= 2:
                data['validasi']['kaprodi_nama'] = clean_rnames[0]
                data['validasi']['ujm_nama'] = clean_rnames[1]

    # Soal UAS
    extracted_soal = extract_soal_uas(text)
    if extracted_soal:
        data['soal_uas'] = extracted_soal

    # If soal not in md, fallback to existing docx in 01_RPS_OBE
    if not data['soal_uas']:
        base_name = os.path.basename(md_path).replace('_FINAL.md', '_OBE.docx')
        alt_docx = os.path.join(r'd:\ANTY GRAVITY\OBE 2026\01_FIK\RPS OBE 2026\RPS_OBE_FINAL\01_RPS_OBE', base_name)
        if os.path.exists(alt_docx):
            try:
                ad = docx.Document(alt_docx)
                in_soal = False
                for p in ad.paragraphs:
                    pt = p.text.strip()
                    if 'CONTOH SOAL' in pt.upper():
                        in_soal = True
                        continue
                    if in_soal:
                        if 'SELAMAT MENGERJAKAN' in pt.upper() or 'LEMBAR VALIDASI' in pt.upper():
                            break
                        if pt and re.match(r'^(?:\d+\.|\bSOAL\b)', pt, re.IGNORECASE):
                            data['soal_uas'].append(pt)
            except Exception as e:
                pass

    return data

def generate_rps_docx(data, template_path=None, output_path=None):
    if template_path is None:
        template_path = DEFAULT_TEMPLATE
    if output_path is None:
        mk_clean = data.get('identitas', {}).get('nama_mk', 'Mata_Kuliah').replace(' ', '_')
        output_path = os.path.join(os.getcwd(), f'RPS_{mk_clean}_OBE.docx')

    doc = docx.Document(template_path)
    
    id_data = data.get('identitas', {})
    otor_data = data.get('otorisasi', {})
    cp_data = data.get('capaian_pembelajaran', {})
    det_data = data.get('detail_mk', {})
    mingguan_data = data.get('mingguan', [])
    rubrik_data = data.get('rubrik', [])
    eval_data = data.get('evaluasi_kelulusan', {})
    soal_data = data.get('soal_uas', [])
    val_data = data.get('validasi', {})

    # ==========================================
    # 1. TABLE 0: Identitas, Otorisasi, CP, Detail MK
    # ==========================================
    t0 = doc.tables[0]
    
    # Row 0: KOP & Kode Dokumen
    kop_lines = [
        "UNIVERSITAS PGRI RONGGOLAWE TUBAN",
        f"FAKULTAS {id_data.get('fakultas', 'KEGURUAN DAN ILMU PENDIDIKAN').upper()}",
        f"PROGRAM STUDI {id_data.get('prodi', 'PENDIDIKAN PANCASILA DAN KEWARGANEGARAAN').upper()}"
    ]
    set_cell_text(t0.rows[0].cells[1], '\n'.join(kop_lines), font_size=11, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(t0.rows[0].cells[7], id_data.get('kode_dokumen', 'PPKn/I/PKN/MKK PPKn'), font_size=10.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)

    # Row 3: Nama MK, Kode, Rumpun, SKS, Semester, Tanggal
    set_cell_text(t0.rows[3].cells[0], id_data.get('nama_mk', ''), font_size=10.5, bold=True)
    set_cell_text(t0.rows[3].cells[2], id_data.get('kode_mk', ''), font_size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(t0.rows[3].cells[3], id_data.get('rumpun_mk', ''), font_size=10.5)
    set_cell_text(t0.rows[3].cells[5], f"{id_data.get('bobot_sks', 2)} SKS", font_size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(t0.rows[3].cells[6], str(id_data.get('semester', 1)), font_size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(t0.rows[3].cells[7], id_data.get('tgl_penyusunan', ''), font_size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER)

    # Row 5: Otorisasi
    set_cell_text(t0.rows[5].cells[2], otor_data.get('pengembang', ''), font_size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(t0.rows[5].cells[4], otor_data.get('koordinator_rmk', ''), font_size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(t0.rows[5].cells[6], otor_data.get('kaprodi', ''), font_size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER)

    # Row 7: CPL-PRODI
    cpl_paragraphs = []
    for cpl in cp_data.get('cpl_prodi', []):
        cpl_paragraphs.append(f"{cpl.get('kode', '')}: {cpl.get('deskripsi', '')}")
    set_cell_text(t0.rows[7].cells[1], '\n\n'.join(cpl_paragraphs), font_size=10.5)

    # Row 9: CPMK
    cpmk_lines = []
    for cpmk in cp_data.get('cpmk', []):
        cpmk_lines.append(f"{cpmk.get('kode', '')}\t{cpmk.get('deskripsi', '')}")
    set_cell_text(t0.rows[9].cells[1], '\n'.join(cpmk_lines), font_size=10.5)

    # Row 11: Sub-CPMK
    subcpmk_lines = []
    for sub in cp_data.get('sub_cpmk', []):
        tag = f" [{sub.get('taksonomi', '')}]" if sub.get('taksonomi') else ""
        subcpmk_lines.append(f"{sub.get('kode', '')}{tag}: {sub.get('deskripsi', '')}")
    set_cell_text(t0.rows[11].cells[1], '\n'.join(subcpmk_lines), font_size=10.5)

    # Row 13: Matriks Korelasi (Nested Table)
    matrix_rows = cp_data.get('matriks_korelasi', [])
    if matrix_rows:
        # Determine CPL list from matrix headers or cpl_prodi
        first_row_cpls = list(matrix_rows[0].get('cpls', {}).keys())
        if first_row_cpls:
            cpl_codes = first_row_cpls
        else:
            cpl_codes = [c.get('kode', '') for c in cp_data.get('cpl_prodi', [])]
        build_nested_matrix(t0.rows[13].cells[1], cpl_codes, matrix_rows)

    # Row 14: Deskripsi Singkat MK
    set_cell_text(t0.rows[14].cells[1], det_data.get('deskripsi', ''), font_size=10.5)

    # Row 15: Bahan Kajian
    bk_text = det_data.get('bahan_kajian', '')
    if isinstance(bk_text, list):
        bk_text = '\n'.join([f"{i+1}. {item}" for i, item in enumerate(bk_text)])
    set_cell_text(t0.rows[15].cells[1], bk_text, font_size=10.5)

    # Row 16: Pustaka Utama
    pu_text = det_data.get('pustaka_utama', '')
    if isinstance(pu_text, list):
        pu_text = '\n'.join(pu_text)
    set_cell_text(t0.rows[16].cells[2], pu_text, font_size=10.5)

    # Row 17: Pustaka Pendukung
    pp_text = det_data.get('pustaka_pendukung', '')
    if isinstance(pp_text, list):
        pp_text = '\n'.join(pp_text)
    set_cell_text(t0.rows[17].cells[2], pp_text, font_size=10.5)

    # Row 18: Dosen Pengampu
    set_cell_text(t0.rows[18].cells[1], det_data.get('dosen_pengampu', ''), font_size=10.5)

    # Row 19: Matakuliah Syarat
    set_cell_text(t0.rows[19].cells[1], det_data.get('matakuliah_syarat', 'Tidak ada'), font_size=10.5)

    # ==========================================
    # 2. TABLE 1: Rencana Pembelajaran 16 Minggu
    # ==========================================
    t1 = doc.tables[1]
    for w_idx, w_data in enumerate(mingguan_data):
        target_row_idx = 3 + w_idx
        if target_row_idx >= len(t1.rows):
            break
        row = t1.rows[target_row_idx]
        mg_no = w_data.get('minggu', w_idx + 1)
        
        if mg_no == 8 or w_data.get('is_uts'):
            set_cell_text(row.cells[0], '8', font_size=10.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
            ind_clean = w_data.get('indikator', 'Evaluasi penguasaan materi perkuliahan minggu 1 s.d. 7')
            ind_clean = re.sub(r'^\s*\*{0,2}Indikator:\*{0,2}\s*', '', ind_clean, flags=re.IGNORECASE).replace('**', '').strip()
            uts_desc = f"UJIAN TENGAH SEMESTER (UTS)\n{ind_clean}"
            set_cell_text(row.cells[1], uts_desc, font_size=10.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
            bobot_clean = str(w_data.get('bobot', 25)).replace('**', '').replace('%', '').strip()
            set_cell_text(row.cells[7], bobot_clean, font_size=10.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
            set_cell_shading(row.cells[0], 'F4F6F7')
            set_cell_shading(row.cells[1], 'F4F6F7')
            set_cell_shading(row.cells[7], 'F4F6F7')
        elif mg_no == 16 or w_data.get('is_uas'):
            set_cell_text(row.cells[0], '16', font_size=10.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
            ind_clean = w_data.get('indikator', 'Evaluasi komprehensif penguasaan capaian pembelajaran mata kuliah')
            ind_clean = re.sub(r'^\s*\*{0,2}Indikator:\*{0,2}\s*', '', ind_clean, flags=re.IGNORECASE).replace('**', '').strip()
            uas_desc = f"UJIAN AKHIR SEMESTER (UAS)\n{ind_clean}"
            set_cell_text(row.cells[1], uas_desc, font_size=10.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
            bobot_clean = str(w_data.get('bobot', 25)).replace('**', '').replace('%', '').strip()
            set_cell_text(row.cells[7], bobot_clean, font_size=10.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
            set_cell_shading(row.cells[0], 'F4F6F7')
            set_cell_shading(row.cells[1], 'F4F6F7')
            set_cell_shading(row.cells[7], 'F4F6F7')
        else:
            set_cell_text(row.cells[0], str(mg_no), font_size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER)
            sub_text = w_data.get('sub_cpmk', '')
            if w_data.get('taksonomi') and f"[{w_data['taksonomi']}]" not in sub_text:
                sub_text = f"{sub_text} [{w_data['taksonomi']}]"
            set_cell_text(row.cells[1], sub_text, font_size=10.5)
            set_cell_text(row.cells[2], w_data.get('indikator', ''), font_size=10.5)
            set_cell_text(row.cells[3], w_data.get('teknik_kriteria', ''), font_size=10.5)
            set_cell_text(row.cells[4], w_data.get('luring', ''), font_size=10.5)
            set_cell_text(row.cells[5], w_data.get('daring', '-'), font_size=10.5)
            set_cell_text(row.cells[6], w_data.get('materi_pustaka', ''), font_size=10.5)
            set_cell_text(row.cells[7], str(w_data.get('bobot', '')).replace('**', '').replace('%', '').strip(), font_size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER)

    # ==========================================
    # 3. TABLE 2: Rubrik Penilaian Holistik
    # ==========================================
    t2 = doc.tables[2]
    for r_idx, r_item in enumerate(rubrik_data):
        target_r = 1 + r_idx
        if target_r >= len(t2.rows):
            break
        row = t2.rows[target_r]
        aspek_clean = str(r_item.get('aspek', '')).replace('**', '').strip()
        bobot_num = str(r_item.get('bobot', '')).replace('**', '').replace('%', '').strip()
        set_cell_text(row.cells[0], f"{aspek_clean} ({bobot_num}%)", font_size=10.5, bold=True)
        set_cell_text(row.cells[1], r_item.get('sangat_baik', ''), font_size=10.5)
        set_cell_text(row.cells[2], r_item.get('baik', ''), font_size=10.5)
        set_cell_text(row.cells[3], r_item.get('cukup', ''), font_size=10.5)
        set_cell_text(row.cells[4], f"{bobot_num}%", font_size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER)

    # ==========================================
    # 4. PARAGRAPHS: Kriteria, Konversi, Soal UAS, Validasi
    # ==========================================
    mk_name = id_data.get('nama_mk', 'Mata Kuliah')
    
    # Identify question paragraph locations
    soal_start_idx = None
    selamat_idx = None
    for i, p in enumerate(doc.paragraphs):
        txt = p.text.strip()
        if txt.startswith("1. ") and soal_start_idx is None:
            soal_start_idx = i
        if 'SELAMAT MENGERJAKAN' in txt.upper() and selamat_idx is None:
            selamat_idx = i

    if soal_start_idx is not None and selamat_idx is not None and soal_data:
        curr_p_idx = soal_start_idx
        for s_idx, s_text in enumerate(soal_data):
            if curr_p_idx < selamat_idx:
                p = doc.paragraphs[curr_p_idx]
                p.text = ''
                s_clean = re.sub(r'^\d+\.\s*', '', s_text.strip())
                add_markdown_runs(p, f"{s_idx + 1}. {s_clean}", font_size=10.5)
                curr_p_idx += 1
            else:
                p = doc.paragraphs[selamat_idx].insert_paragraph_before()
                s_clean = re.sub(r'^\d+\.\s*', '', s_text.strip())
                add_markdown_runs(p, f"{s_idx + 1}. {s_clean}", font_size=10.5)
        # Clear unused placeholder questions
        while curr_p_idx < selamat_idx:
            if any(doc.paragraphs[curr_p_idx].text.strip().startswith(f"{n}. ") for n in range(1, 10)):
                doc.paragraphs[curr_p_idx].text = ''
            curr_p_idx += 1

    for p in doc.paragraphs:
        txt = p.text
        if 'Mahasiswa dinyatakan lulus mata kuliah' in txt:
            p.text = f"Mahasiswa dinyatakan lulus mata kuliah {mk_name} apabila memperoleh nilai akhir minimal 60 (C) dengan komponen penilaian sebagai berikut:"
            p.runs[0].font.name = 'Cambria'
            p.runs[0].font.size = Pt(10.5)
        elif 'Keaktifan & Partisipasi:' in txt and 'keaktifan' in eval_data:
            p.text = f"Keaktifan & Partisipasi: {eval_data.get('keaktifan', 10)}%"
            p.runs[0].font.name = 'Cambria'
            p.runs[0].font.size = Pt(10.5)
        elif 'Tugas & Kuis:' in txt and 'tugas' in eval_data:
            p.text = f"Tugas & Kuis: {eval_data.get('tugas', 15)}%"
            p.runs[0].font.name = 'Cambria'
            p.runs[0].font.size = Pt(10.5)
        elif 'UTS:' in txt and 'uts' in eval_data:
            p.text = f"UTS: {eval_data.get('uts', 25)}%"
            p.runs[0].font.name = 'Cambria'
            p.runs[0].font.size = Pt(10.5)
        elif 'Proyek Akhir:' in txt and 'proyek' in eval_data:
            p.text = f"Proyek Akhir: {eval_data.get('proyek', 15)}%"
            p.runs[0].font.name = 'Cambria'
            p.runs[0].font.size = Pt(10.5)
        elif 'UAS:' in txt and 'uas' in eval_data:
            p.text = f"UAS: {eval_data.get('uas', 35)}%"
            p.runs[0].font.name = 'Cambria'
            p.runs[0].font.size = Pt(10.5)
        elif 'Rencana Pembelajaran Semester (RPS) mata kuliah' in txt:
            p.text = f"Rencana Pembelajaran Semester (RPS) mata kuliah {mk_name} ini telah divalidasi dan dinyatakan memenuhi standar mutu kurikulum OBE Universitas PGRI Ronggolawe (UNIROW) Tuban."
            p.runs[0].font.name = 'Cambria'
            p.runs[0].font.size = Pt(10.5)
        elif 'Divalidasi di Tuban, pada tanggal:' in txt:
            p.text = f"Divalidasi di Tuban, pada tanggal: {val_data.get('tgl_validasi', id_data.get('tgl_penyusunan', '04 Februari 2026'))}"
            p.runs[0].font.name = 'Cambria'
            p.runs[0].font.size = Pt(10.5)

    # ==========================================
    # 5. TABLE 3: Lembar Validasi Tanda Tangan
    # ==========================================
    t3 = doc.tables[3]
    prodi_name = id_data.get('prodi', 'PPKn')
    
    set_cell_text(t3.rows[0].cells[0], f"Menyetujui,\nKetua Program Studi {prodi_name}", font_size=10.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(t3.rows[0].cells[1], f"Mengetahui,\nUnit Jaminan Mutu (UJM) Prodi {prodi_name}", font_size=10.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_shading(t3.rows[0].cells[0], 'FAFAFA')
    set_cell_shading(t3.rows[0].cells[1], 'FAFAFA')

    kaprodi_txt = f"\n\n\n    {val_data.get('kaprodi_nama', otor_data.get('kaprodi', 'Mario Fahmi Syahrial, M.Pd.'))}\n    NIDN. {val_data.get('kaprodi_nidn', '....................')}"
    ujm_txt = f"\n\n\n    {val_data.get('ujm_nama', '[Tim Penjamin Mutu Prodi PPKn]')}\n    NIDN. {val_data.get('ujm_nidn', '....................')}"
    set_cell_text(t3.rows[1].cells[0], kaprodi_txt, font_size=10.5, bold=True)
    set_cell_text(t3.rows[1].cells[1], ujm_txt, font_size=10.5, bold=True)
    set_cell_shading(t3.rows[1].cells[0], 'FAFAFA')
    set_cell_shading(t3.rows[1].cells[1], 'FAFAFA')

    # Document-wide asterisk cleaning pass (zero tolerance)
    for p in doc.paragraphs:
        for r in p.runs:
            if '**' in r.text or '__' in r.text:
                r.text = r.text.replace('**', '').replace('__', '')
                
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        if '**' in r.text or '__' in r.text:
                            r.text = r.text.replace('**', '').replace('__', '')
                for ntbl in cell.tables:
                    for nrow in ntbl.rows:
                        for ncell in nrow.cells:
                            for np in ncell.paragraphs:
                                for nr in np.runs:
                                    if '**' in nr.text or '__' in nr.text:
                                        nr.text = nr.text.replace('**', '').replace('__', '')

    doc.save(output_path)
    print(f"[OK] RPS berhasil dibuat: {output_path}")
    return output_path

def main():
    parser = argparse.ArgumentParser(description="Generator Dokumen RPS UNIROW OBE 2026 Presisi")
    parser.add_argument('--md', help="Path file RPS format Markdown (.md)")
    parser.add_argument('--data', help="Path file data JSON RPS (.json)")
    parser.add_argument('--template', default=DEFAULT_TEMPLATE, help="Path master template DOCX")
    parser.add_argument('--output', help="Path output DOCX RPS")
    args = parser.parse_args()

    data = None
    if args.md:
        data = parse_rps_markdown(args.md)
    elif args.data:
        with open(args.data, 'r', encoding='utf-8') as f:
            data = json.load(f)
    else:
        parser.print_help()
        return

    generate_rps_docx(data, args.template, args.output)

if __name__ == '__main__':
    main()
